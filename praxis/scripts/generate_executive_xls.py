#!/usr/bin/env python3
"""
Usage: python generate_executive_xls.py --manifest <path> [--output <path>] [--master <path>]
Generates A-OPS executive spreadsheet with 5 tabs using openpyxl.
"""
import argparse
import datetime
import os
import sys
import yaml

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# Design system colors
COLOR_PRIMARY = "1B2A4A"
COLOR_ACCENT = "2E7D9B"
COLOR_SURFACE = "F5F7FA"
COLOR_NEUTRAL = "5C6B7A"
COLOR_WHITE = "FFFFFF"
COLOR_DIVIDER = "D0D7E0"

HEADER_FONT = Font(name="Calibri", bold=True, color=COLOR_WHITE, size=10)
BODY_FONT = Font(name="Calibri", size=10)
HEADER_FILL = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
ACCENT_FILL = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type="solid")
ALT_FILL = PatternFill(start_color=COLOR_SURFACE, end_color=COLOR_SURFACE, fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color=COLOR_DIVIDER),
    right=Side(style="thin", color=COLOR_DIVIDER),
    top=Side(style="thin", color=COLOR_DIVIDER),
    bottom=Side(style="thin", color=COLOR_DIVIDER),
)


def style_header_row(ws, row: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 20


def style_data_row(ws, row: int, num_cols: int, alternate: bool = False):
    fill = ALT_FILL if alternate else PatternFill(fill_type=None)
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BODY_FONT
        if alternate:
            cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 18


def set_column_widths(ws, widths: list):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(width, 15)


def build_resumo_executivo(wb, manifest: dict):
    ws = wb.active
    ws.title = "Resumo Executivo"
    ws.sheet_properties.tabColor = COLOR_ACCENT

    headers = ["Campo", "Valor"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    client = manifest.get("client_identity", {})
    branding = client.get("branding", {})
    gates_passed = manifest.get("gates_passed", [])
    artifacts = manifest.get("artifacts_produced", [])

    rows = [
        ("Case ID", manifest.get("case_id", "—")),
        ("Consultor", branding.get("consultant_display_name", manifest.get("consultant_id", "—"))),
        ("Empresa do Consultor", branding.get("consultant_company", "—")),
        ("Cliente", client.get("company_name", "—")),
        ("Segmento", client.get("segment", "—")),
        ("Tamanho da Equipe", client.get("team_size", "—")),
        ("Receita Anual (faixa)", client.get("annual_revenue_range", "—")),
        ("Decisores", ", ".join(client.get("decision_makers", []))),
        ("Cenário", manifest.get("scenario_selected", "—")),
        ("Fase Atual", f"Fase {manifest.get('current_phase', '?')} de 6"),
        ("Criado em", manifest.get("created_at", "—")),
        ("Atualizado em", manifest.get("updated_at", "—")),
        ("Gates Aprovados", ", ".join(g.get("gate", "") for g in gates_passed)),
        ("Gates Pendentes", ", ".join(manifest.get("gates_pending", []))),
        ("Total de Artefatos", str(len(artifacts))),
    ]

    for i, (campo, valor) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=campo)
        ws.cell(row=i, column=2, value=str(valor))
        style_data_row(ws, i, 2, alternate=(i % 2 == 0))

    set_column_widths(ws, [30, 60])


def build_diagnostico(wb, manifest: dict):
    ws = wb.create_sheet("Diagnóstico")
    ws.sheet_properties.tabColor = COLOR_NEUTRAL

    headers = ["#", "Problema / Causa", "Rótulo Epistêmico", "Gravidade (G)", "Urgência (U)", "Tendência (T)", "Score GUT", "Prioridade"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    # Placeholder rows — in production, parsed from A-07/A-08/A-09
    placeholder_rows = [
        (1, "Problema central identificado no briefing", "[FATO]", 5, 5, 5, 125, "Alta"),
        (2, "Causa raiz primária (a definir na Fase 3)", "[HIPÓTESE]", 4, 4, 3, 48, "Alta"),
        (3, "Causa raiz secundária (a definir na Fase 3)", "[HIPÓTESE]", 3, 3, 3, 27, "Média"),
    ]

    for i, row in enumerate(placeholder_rows, 2):
        for j, val in enumerate(row, 1):
            ws.cell(row=i, column=j, value=val)
        style_data_row(ws, i, len(headers), alternate=(i % 2 == 0))

    set_column_widths(ws, [5, 50, 20, 12, 12, 12, 12, 12])


def build_plano_acao(wb, manifest: dict):
    ws = wb.create_sheet("Plano de Ação")
    ws.sheet_properties.tabColor = COLOR_ACCENT

    headers = ["#", "O quê (What)", "Por quê (Why)", "Quem (Who)", "Quando (When)", "Onde (Where)", "Como (How)", "Quanto (How Much)", "Status", "KPI"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    # Placeholder row
    placeholder = [1, "Ação a ser definida na Fase 3", "Impacto no problema central [INFERÊNCIA]",
                   "Responsável a definir", "Prazo a definir", "—", "Como a definir", "R$ a definir", "Pendente", "KPI a definir"]
    ws.append(placeholder)
    style_data_row(ws, 2, len(headers))

    set_column_widths(ws, [5, 40, 40, 20, 15, 15, 35, 15, 12, 25])


def build_simulacao(wb, manifest: dict):
    ws = wb.create_sheet("Simulação")
    ws.sheet_properties.tabColor = COLOR_NEUTRAL

    headers = ["Métrica", "Valor Atual", "Cenário Pessimista [HIPÓTESE]", "Cenário Conservador [HIPÓTESE]", "Cenário Otimista [HIPÓTESE]", "Premissa-Chave", "Validado?"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    metrics = [
        ("Receita Mensal", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("Número de Clientes", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("Ticket Médio", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("Churn Mensal (%)", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("CAC (R$)", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("LTV (R$)", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
        ("Margem Bruta (%)", "—", "—", "—", "—", "A definir na Fase 4", "Não"),
    ]

    for i, row in enumerate(metrics, 2):
        for j, val in enumerate(row, 1):
            ws.cell(row=i, column=j, value=val)
        style_data_row(ws, i, len(headers), alternate=(i % 2 == 0))

    set_column_widths(ws, [25, 18, 28, 28, 28, 40, 12])


def build_proximos_passos(wb, manifest: dict):
    ws = wb.create_sheet("Próximos Passos")
    ws.sheet_properties.tabColor = COLOR_ACCENT

    headers = ["#", "Ação Imediata", "Responsável", "Prazo", "KPI de Acompanhamento", "Dependências", "Status"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    placeholder = [1, "Próximos passos a definir após Gate G2 (diagnóstico aprovado)", "—", "—", "—", "—", "Pendente"]
    ws.append(placeholder)
    style_data_row(ws, 2, len(headers))

    set_column_widths(ws, [5, 50, 20, 15, 35, 30, 12])


def main():
    parser = argparse.ArgumentParser(description="Generate A-OPS executive spreadsheet")
    parser.add_argument("--manifest", default="manifest.yaml")
    parser.add_argument("--output", default=None)
    parser.add_argument("--master", default=None, help="Path to A-MASTER for data extraction")
    args = parser.parse_args()

    try:
        with open(args.manifest, encoding="utf-8") as f:
            manifest = yaml.safe_load(f)
    except FileNotFoundError:
        print(f"ERROR: manifest not found: {args.manifest}", file=sys.stderr)
        sys.exit(1)

    case_id = manifest.get("case_id", "unknown")
    client_name = manifest.get("client_identity", {}).get("company_name", "cliente")
    safe_client = "".join(c if c.isalnum() or c in "-_" else "_" for c in client_name)
    output_path = args.output or f"A-OPS_{safe_client}_{case_id[:16]}.xlsx"

    # Override colors from branding if configured
    branding = manifest.get("client_identity", {}).get("branding", {})
    if branding.get("primary_color"):
        hex_color = branding["primary_color"].lstrip("#")
        HEADER_FILL.start_color.rgb = hex_color
        HEADER_FILL.end_color.rgb = hex_color

    wb = openpyxl.Workbook()
    build_resumo_executivo(wb, manifest)
    build_diagnostico(wb, manifest)
    build_plano_acao(wb, manifest)
    build_simulacao(wb, manifest)
    build_proximos_passos(wb, manifest)

    wb.save(output_path)
    print(f"A-OPS generated: {output_path}")
    print(f"Tabs: Resumo Executivo, Diagnóstico, Plano de Ação, Simulação, Próximos Passos")


if __name__ == "__main__":
    main()
